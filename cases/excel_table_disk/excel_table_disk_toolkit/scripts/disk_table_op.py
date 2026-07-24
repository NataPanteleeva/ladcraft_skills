async def handler(state, params):
    warnings = []
    params = params if isinstance(params, dict) else {}
    user_env = read_user_env(state if isinstance(state, dict) else {})
    skill_storage = resolve_skill_storage(state if isinstance(state, dict) else {})
    base_url = resolve_base_url(state if isinstance(state, dict) else {}, params, skill_storage)
    login = pick_string(params.get("login"), user_env.get("R7_DISK_LOGIN"))
    password = pick_string(params.get("password"), user_env.get("R7_DISK_PASSWORD"))
    # __PUBLISH_ENV_FALLBACK__

    if not base_url:
        return {
            "ok": False,
            "error": "Не задан R7_DISK_BASE_URL в настройках навыка.",
            "agent_message": "Навык не настроен (R7_DISK_*). Обратитесь к администратору.",
        }

    operation = pick_string(params.get("operation"), "profile").lower()
    doc_a = resolve_positive_id(params.get("document_id") or params.get("document_id_a"))
    doc_b = resolve_positive_id(params.get("document_id_b"))
    directory_id = resolve_positive_id(params.get("directory_id") or params.get("output_directory_id"))
    sheet_a = pick_string(params.get("sheetA") or params.get("sourceSheet"))
    sheet_b = pick_string(params.get("sheetB"))
    output_name = ensure_xlsx_extension(
        pick_string(params.get("output_name"), "результат_таблицы.xlsx")
    )
    conflict_policy = pick_string(params.get("conflict_policy"), "overwrite").lower() or "overwrite"

    auth_result = await ensure_auth_token(base_url, login, password, skill_storage, params.get("auth_token"))
    if not auth_result.get("ok"):
        return {
            "ok": False,
            "error": auth_result.get("error", "Ошибка авторизации."),
            "agent_message": "Не удалось войти на Р7-Диск. Проверьте R7_DISK_*.",
        }
    auth_token = auth_result["auth_token"]

    if doc_a is None:
        return {"ok": False, "error": "document_id / document_id_a обязателен", "agent_message": "Укажите document_id файла на диске."}

    dl_a = await download_document_bytes(base_url, auth_token, doc_a)
    if not dl_a.get("ok"):
        return {"ok": False, "error": "Не скачан A: " + str(dl_a.get("error")), "agent_message": str(dl_a.get("error"))}

    from io import BytesIO
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
    import math
    import re
    from collections import Counter

    def cell_key(value):
        if value is None:
            return ""
        if isinstance(value, float) and math.isfinite(value) and value == int(value):
            return str(int(value))
        return str(value).strip()

    def to_number(value):
        if isinstance(value, (int, float)):
            n = float(value)
            return n if math.isfinite(n) else float("nan")
        if value is None or value == "":
            return float("nan")
        t = str(value).strip().replace(" ", "").replace(",", ".")
        try:
            n = float(t)
            return n if math.isfinite(n) else float("nan")
        except Exception:
            return float("nan")

    def round_smart(value):
        if not isinstance(value, (int, float)) or not math.isfinite(float(value)):
            return value
        if float(value).is_integer():
            return int(value)
        return round(float(value), 2)

    def normalize_header(h):
        return re.sub(r"\s+", " ", str(h or "").strip())

    def sheet_to_records(wb, sheet_name):
        names = list(wb.sheetnames)
        name = sheet_name if sheet_name and sheet_name in wb.sheetnames else (names[0] if names else "")
        if not name:
            raise ValueError("Лист не найден")
        ws = wb[name]
        rows_iter = list(ws.iter_rows(values_only=True))
        if not rows_iter:
            return {"sheetName": name, "headers": [], "rows": []}
        headers = [normalize_header(h) for h in rows_iter[0]]
        headers = [h if h else ("col_" + str(i + 1)) for i, h in enumerate(headers)]
        rows = []
        for raw in rows_iter[1:]:
            if raw is None or all(v is None or v == "" for v in raw):
                continue
            item = {}
            for i, header in enumerate(headers):
                value = raw[i] if i < len(raw) else None
                if hasattr(value, "isoformat"):
                    value = value.isoformat()
                item[header] = value
            rows.append(item)
        return {"sheetName": name, "headers": headers, "rows": rows}

    def records_to_aoa(headers, rows):
        aoa = [list(headers)]
        for row in rows:
            aoa.append([row.get(h) for h in headers])
        return aoa

    def write_sheets_bytes(sheets_aoa):
        wb = Workbook()
        wb.remove(wb.active)
        for item in sheets_aoa:
            name = str(item.get("name") or "Результат")[:31]
            aoa = item.get("aoa") or []
            ws = wb.create_sheet(title=name)
            for r_idx, row in enumerate(aoa, start=1):
                for c_idx, value in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=value)
            width = max(1, len(aoa[0]) if aoa else 1)
            for col in range(1, width + 1):
                ws.column_dimensions[get_column_letter(col)].width = 16
        out = BytesIO()
        wb.save(out)
        return out.getvalue()

    wb_a = load_workbook(BytesIO(dl_a["bytes"]), data_only=True)
    parsed_a = sheet_to_records(wb_a, sheet_a)
    sheets_out = []
    summary = {"operation": operation, "sheetA": parsed_a["sheetName"], "rowCount": len(parsed_a["rows"])}

    if operation == "compare":
        if doc_b is None:
            return {"ok": False, "error": "для compare нужен document_id_b", "agent_message": "Укажите второй файл (document_id_b)."}
        dl_b = await download_document_bytes(base_url, auth_token, doc_b)
        if not dl_b.get("ok"):
            return {"ok": False, "error": str(dl_b.get("error")), "agent_message": str(dl_b.get("error"))}
        wb_b = load_workbook(BytesIO(dl_b["bytes"]), data_only=True)
        parsed_b = sheet_to_records(wb_b, sheet_b)
        key_fields = params.get("keyFields") if isinstance(params.get("keyFields"), list) else []
        key_fields = [str(x).strip() for x in key_fields if str(x).strip()]
        if not key_fields:
            common = [h for h in parsed_a["headers"] if h in parsed_b["headers"]]
            if not common:
                return {"ok": False, "error": "Нет общих колонок для ключа", "agent_message": "Укажите keyFields."}
            key_fields = [common[0]]
            warnings.append("Ключ авто: " + key_fields[0])

        def index_rows(parsed):
            out = {}
            for row in parsed["rows"]:
                key = "\u0001".join(cell_key(row.get(f)) for f in key_fields)
                if key not in out:
                    out[key] = row
            return out

        map_a = index_rows(parsed_a)
        map_b = index_rows(parsed_b)
        only_a = sorted(set(map_a) - set(map_b))
        only_b = sorted(set(map_b) - set(map_a))
        both = sorted(set(map_a) & set(map_b))
        sheets_out.append({"name": "Только_A", "aoa": records_to_aoa(parsed_a["headers"], [map_a[k] for k in only_a])})
        sheets_out.append({"name": "Только_B", "aoa": records_to_aoa(parsed_b["headers"], [map_b[k] for k in only_b])})
        sheets_out.append({"name": "Совпадают", "aoa": records_to_aoa(parsed_a["headers"], [map_a[k] for k in both])})
        summary.update({"keyFields": key_fields, "onlyA": len(only_a), "onlyB": len(only_b), "matched": len(both), "sheetB": parsed_b["sheetName"]})
        output_name = ensure_xlsx_extension(pick_string(params.get("output_name"), "сверка_диск.xlsx"))

    elif operation == "profile":
        aoa = [["Колонка", "nonEmpty", "nullPct", "unique", "top"]]
        n = len(parsed_a["rows"])
        for header in parsed_a["headers"]:
            values = [row.get(header) for row in parsed_a["rows"]]
            non_empty = [v for v in values if v is not None and v != ""]
            null_pct = round(100.0 * (n - len(non_empty)) / n, 1) if n else 0.0
            uniq = len({cell_key(v) for v in non_empty})
            top = Counter(cell_key(v) for v in non_empty).most_common(5)
            top_str = "; ".join(k + "(" + str(c) + ")" for k, c in top)
            aoa.append([header, len(non_empty), null_pct, uniq, top_str])
        sheets_out.append({"name": "Профиль", "aoa": aoa})
        output_name = ensure_xlsx_extension(pick_string(params.get("output_name"), "профиль_диск.xlsx"))

    elif operation == "filter":
        filters = params.get("filters") if isinstance(params.get("filters"), list) else []
        rows = list(parsed_a["rows"])
        for rule in filters:
            if not isinstance(rule, dict):
                continue
            field = str(rule.get("field") or "").strip()
            op = str(rule.get("op") or "eq").lower()
            expected = rule.get("value")
            kept = []
            for row in rows:
                actual = row.get(field)
                ok = True
                if op == "eq":
                    ok = cell_key(actual) == cell_key(expected)
                elif op == "contains":
                    ok = str(expected or "").lower() in str(actual or "").lower()
                elif op == "gt":
                    ok = to_number(actual) > to_number(expected)
                elif op == "lt":
                    ok = to_number(actual) < to_number(expected)
                if ok:
                    kept.append(row)
            rows = kept
        sheets_out.append({"name": "Фильтр", "aoa": records_to_aoa(parsed_a["headers"], rows)})
        summary["rowCountAfter"] = len(rows)
        output_name = ensure_xlsx_extension(pick_string(params.get("output_name"), "фильтр_диск.xlsx"))

    elif operation == "pivot":
        row_fields = [str(x).strip() for x in (params.get("rowFields") or []) if str(x).strip()]
        column_fields = [str(x).strip() for x in (params.get("columnFields") or []) if str(x).strip()]
        value_field = pick_string(params.get("valueField"))
        aggregation = pick_string(params.get("aggregation"), "sum").lower() or "sum"
        if not row_fields or not value_field:
            return {"ok": False, "error": "pivot: нужны rowFields и valueField", "agent_message": "Укажите поля сводной."}
        groups = {}
        col_set = {}
        column_field = column_fields[0] if column_fields else ""
        for row in parsed_a["rows"]:
            rk = "\u0001".join(cell_key(row.get(f)) for f in row_fields)
            ck = cell_key(row.get(column_field)) if column_field else "__value__"
            if column_field:
                col_set[ck] = True
            if rk not in groups:
                groups[rk] = {"parts": [cell_key(row.get(f)) for f in row_fields], "cols": {}}
            groups[rk]["cols"].setdefault(ck, [])
            if aggregation == "count":
                groups[rk]["cols"][ck].append(1.0)
            else:
                n = to_number(row.get(value_field))
                if math.isfinite(n):
                    groups[rk]["cols"][ck].append(n)
        col_keys = sorted(col_set.keys()) if column_field else ["__value__"]
        header = list(row_fields) + (col_keys if column_field else [value_field + " (" + aggregation + ")"])
        body = []
        for rk in sorted(groups.keys()):
            g = groups[rk]
            line = list(g["parts"])
            for ck in col_keys:
                vals = g["cols"].get(ck) or []
                if not vals:
                    line.append(None)
                elif aggregation == "avg":
                    line.append(round_smart(sum(vals) / len(vals)))
                elif aggregation == "min":
                    line.append(round_smart(min(vals)))
                elif aggregation == "max":
                    line.append(round_smart(max(vals)))
                else:
                    line.append(round_smart(sum(vals)))
            body.append(line)
        sheets_out.append({"name": "Сводная", "aoa": [header] + body})
        summary["pivotRows"] = len(body)
        output_name = ensure_xlsx_extension(pick_string(params.get("output_name"), "сводная_диск.xlsx"))

    elif operation == "top_n":
        category_field = pick_string(params.get("categoryField"))
        value_field = pick_string(params.get("valueField"))
        top_n = int(params.get("topN") or 10)
        if not category_field or not value_field:
            return {"ok": False, "error": "top_n: categoryField и valueField", "agent_message": "Укажите categoryField и valueField."}
        agg = {}
        for row in parsed_a["rows"]:
            k = cell_key(row.get(category_field))
            agg.setdefault(k, {"label": row.get(category_field), "sum": 0.0})
            n = to_number(row.get(value_field))
            if math.isfinite(n):
                agg[k]["sum"] += n
        ranked = sorted(agg.values(), key=lambda x: -x["sum"])[: max(1, min(50, top_n))]
        aoa = [[category_field, value_field]] + [[r["label"], round_smart(r["sum"])] for r in ranked]
        sheets_out.append({"name": "Топ", "aoa": aoa})
        output_name = ensure_xlsx_extension(pick_string(params.get("output_name"), "топ_диск.xlsx"))

    else:
        return {
            "ok": False,
            "error": "Неизвестная operation: " + operation,
            "agent_message": "operation: profile | filter | pivot | top_n | compare",
        }

    output_bytes = write_sheets_bytes(sheets_out)

    if directory_id is None:
        directory_id = await resolve_personal_root_directory_id(base_url, auth_token, skill_storage, user_env)
    if directory_id is None:
        return {"ok": False, "error": "Не удалось определить папку для сохранения", "agent_message": "Укажите directory_id."}

    final_name_result = await resolve_output_name_by_policy(
        base_url, auth_token, directory_id, output_name, conflict_policy
    )
    if not final_name_result.get("ok"):
        return {"ok": False, "error": final_name_result.get("error"), "agent_message": final_name_result.get("error")}

    uploaded = await upload_replacing_document(
        base_url,
        auth_token,
        directory_id,
        final_name_result["name"],
        output_bytes,
        XLSX_MIME,
    )
    if not uploaded.get("ok"):
        return {"ok": False, "error": uploaded.get("error"), "agent_message": uploaded.get("error")}

    msg = (
        "Готово: «"
        + final_name_result["name"]
        + "» на Р7-Диске (directory_id="
        + str(directory_id)
        + ", document_id="
        + str(uploaded.get("document_id"))
        + ")."
    )
    return {
        "ok": True,
        "operation": operation,
        "directory_id": directory_id,
        "output_name": final_name_result["name"],
        "output_document_id": uploaded.get("document_id"),
        "output_size_bytes": len(output_bytes),
        "summary": summary,
        "warnings": warnings,
        "agent_message": msg,
    }
