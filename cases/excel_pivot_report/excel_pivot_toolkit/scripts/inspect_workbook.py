async def handler(state, params):
    params = params if isinstance(params, dict) else {}
    vfs = get_vfs(state)
    if not vfs:
        return {"ok": False, "error": "VFS недоступен"}

    file_path = get_string(params, "path").strip()
    if not file_path:
        return {"ok": False, "error": "path обязателен"}

    sheet = get_string(params, "sheet").strip()
    sample_rows = 5
    raw_sample = params.get("sampleRows")
    if isinstance(raw_sample, (int, float)) and math.isfinite(float(raw_sample)):
        sample_rows = max(1, min(20, int(round(float(raw_sample)))))

    try:
        from openpyxl import load_workbook

        buffer = await read_workbook_bytes(vfs, file_path)
        wb = load_workbook(io.BytesIO(buffer), data_only=True)
        parsed = sheet_to_records(wb, sheet)
        columns = classify_columns(parsed["headers"], parsed["rows"], 50)
        suggested = suggest_pivot(columns)
        sample = parsed["rows"][:sample_rows]
        return {
            "ok": True,
            "path": file_path,
            "sheets": list(wb.sheetnames),
            "sheet": parsed["sheetName"],
            "headers": parsed["headers"],
            "rowCount": len(parsed["rows"]),
            "sample": sample,
            "columns": columns,
            "suggestedPivot": suggested,
        }
    except Exception as err:
        return {"ok": False, "error": str(err), "path": file_path}
