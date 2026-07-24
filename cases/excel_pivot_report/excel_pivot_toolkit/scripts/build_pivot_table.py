async def handler(state, params):
    params = params if isinstance(params, dict) else {}
    vfs = get_vfs(state)
    if not vfs:
        return {"ok": False, "error": "VFS недоступен"}

    source_path = get_string(params, "sourcePath").strip()
    target_path = normalize_result_path(
        get_string(params, "targetPath").strip(),
        source_path,
        "pivot",
    )
    source_sheet = get_string(params, "sourceSheet").strip()
    target_sheet = get_string(params, "targetSheet").strip() or "Сводная"
    row_fields = [str(item or "").strip() for item in get_array(params, "rowFields") if str(item or "").strip()]
    column_fields = [str(item or "").strip() for item in get_array(params, "columnFields") if str(item or "").strip()]
    value_field = get_string(params, "valueField").strip()
    aggregation = get_string(params, "aggregation").strip().lower() or "sum"
    if aggregation not in ("sum", "count", "avg", "min", "max"):
        aggregation = "sum"

    preview_rows = 10
    raw_preview = params.get("previewRows")
    if isinstance(raw_preview, (int, float)) and math.isfinite(float(raw_preview)):
        preview_rows = max(1, min(30, int(round(float(raw_preview)))))

    if not source_path:
        return {"ok": False, "error": "sourcePath обязателен"}
    if not row_fields:
        return {"ok": False, "error": "rowFields обязателен (минимум одно поле)"}
    if not value_field:
        return {"ok": False, "error": "valueField обязателен"}

    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.utils import get_column_letter

        source_buffer = await read_workbook_bytes(vfs, source_path)
        source_wb = load_workbook(io.BytesIO(source_buffer), data_only=True)
        parsed = sheet_to_records(source_wb, source_sheet)

        header_set = {h: True for h in parsed["headers"]}
        missing = []
        for field in row_fields:
            if field not in header_set:
                missing.append(field)
        for field in column_fields:
            if field not in header_set:
                missing.append(field)
        if value_field not in header_set:
            missing.append(value_field)
        if missing:
            return {
                "ok": False,
                "error": "Неизвестные колонки: " + ", ".join(missing) + ". Доступны: " + ", ".join(parsed["headers"]),
            }

        pivot = build_pivot_aoa(
            parsed["rows"],
            {
                "rowFields": row_fields,
                "columnFields": column_fields,
                "valueField": value_field,
                "aggregation": aggregation,
            },
        )

        created_new = False
        if target_path == source_path:
            # Reload without data_only to preserve formulas/styles if any; for smoke we rewrite from buffer.
            target_wb = load_workbook(io.BytesIO(source_buffer), data_only=False)
        else:
            exists = False
            if callable(getattr(vfs, "exists", None)):
                try:
                    exists = bool(await vfs.exists(target_path))
                except Exception:
                    exists = False
            if exists:
                target_buffer = await read_workbook_bytes(vfs, target_path)
                target_wb = load_workbook(io.BytesIO(target_buffer), data_only=False)
            else:
                target_wb = Workbook()
                created_new = True
                default = target_wb.active
                target_wb.remove(default)

        if target_sheet in target_wb.sheetnames:
            del target_wb[target_sheet]
        ws = target_wb.create_sheet(title=target_sheet)

        for r_idx, row in enumerate(pivot["aoa"], start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        for col in range(1, pivot["columnCount"] + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16

        out = io.BytesIO()
        target_wb.save(out)
        payload = out.getvalue()
        write_mode = await write_workbook_bytes(vfs, target_path, payload)

        stored = None
        if callable(getattr(vfs, "getFileMetadata", None)):
            try:
                meta = await vfs.getFileMetadata(target_path)
                if isinstance(meta, dict):
                    stored = meta.get("size_bytes") or meta.get("size")
            except Exception:
                stored = None

        return analytics_done(
            target_path,
            "Сводная «"
            + target_sheet
            + "»: "
            + str(pivot["rowCount"])
            + " строк, "
            + aggregation
            + " по «"
            + value_field
            + "».",
            sourcePath=source_path,
            targetSheet=target_sheet,
            writeMode=write_mode + ("+created" if created_new else "+updated"),
            expectedBytes=len(payload),
            storedBytes=stored,
            pivotRows=pivot["rowCount"],
            pivotColumns=pivot["columnCount"],
            aggregation=pivot["aggregation"],
            valueField=pivot["valueField"],
            rowFields=row_fields,
            columnFields=column_fields[:1],
            preview=pivot["aoa"][: preview_rows + 1],
        )
    except Exception as err:
        return {
            "ok": False,
            "error": str(err),
            "sourcePath": source_path,
            "targetPath": target_path,
            "targetSheet": target_sheet,
        }
