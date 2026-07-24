async def handler(state, params):
    params = params if isinstance(params, dict) else {}
    vfs = get_vfs(state)
    if not vfs:
        return {"ok": False, "error": "VFS недоступен", "needsUser": False, "options": []}

    workbook_path = get_string(params, "workbookPath").strip()
    last_result = get_string(params, "lastResultPath").strip()
    user_choice = get_string(params, "userChoice").strip()
    hint = get_string(params, "hint").strip()

    skip_names = {
        "history.md",
        ".tmp",
        "_excel_pivot_out.xlsx",
    }

    def norm_path(p):
        s = str(p or "").strip().replace("\\", "/")
        if s.startswith("~/"):
            s = s[1:]
        if not s.startswith("/") and s.startswith("session/"):
            s = "/" + s
        return s

    def is_excel_path(p):
        lower = p.lower()
        return lower.endswith(".xlsx") or lower.endswith(".xlsm") or lower.endswith(".xls")

    def short_label(path, sheet):
        name = path.rsplit("/", 1)[-1] if path else "файл"
        if name.lower().startswith("r7-"):
            name = "Исходная книга"
        return name + " · " + str(sheet)

    candidate_paths = []
    for raw in (workbook_path, last_result):
        p = norm_path(raw)
        if p and is_excel_path(p) and p not in candidate_paths:
            candidate_paths.append(p)

    discovered = []
    for root in ("/session/r7", "/session"):
        await list_excel_in_dir(vfs, root, discovered, 0)
    for item in discovered:
        p = norm_path(item.get("path"))
        name = str(item.get("name") or "").strip()
        if not p or not is_excel_path(p):
            continue
        if name.lower() in skip_names:
            continue
        if p not in candidate_paths:
            candidate_paths.append(p)

    options = []
    seen = set()
    idx = 0
    for path in candidate_paths:
        try:
            from openpyxl import load_workbook

            buffer = await read_workbook_bytes(vfs, path)
            wb = load_workbook(io.BytesIO(buffer), read_only=True, data_only=True)
            sheets = list(wb.sheetnames or [])
            try:
                wb.close()
            except Exception:
                pass
        except Exception:
            continue
        if not sheets:
            continue
        for sheet in sheets:
            key = path + "::" + sheet
            if key in seen:
                continue
            seen.add(key)
            idx += 1
            opt_id = "s" + str(idx)
            options.append(
                {
                    "id": opt_id,
                    "label": short_label(path, sheet),
                    "sourcePath": path,
                    "sourceSheet": sheet,
                }
            )

    if not options:
        return {
            "ok": False,
            "error": "Нет доступных .xlsx с листами в session VFS",
            "needsUser": False,
            "options": [],
        }

    def resolve_opt(opt):
        return {
            "ok": True,
            "needsUser": False,
            "sourcePath": opt["sourcePath"],
            "sourceSheet": opt["sourceSheet"],
            "label": opt["label"],
            "options": options,
        }

    def is_open_workbook(path):
        p = norm_path(path).lower()
        return "/session/r7/" in p or p.startswith("/session/r7/")

    def is_result_artifact(path):
        p = norm_path(path).lower()
        name = p.rsplit("/", 1)[-1]
        if is_open_workbook(p):
            return False
        # Agent deliverables: KPI / pivot / filter exports — not the raw open book.
        markers = ("_kpi", "kpi_", "premium_", "_sorted", "_filtered", "_pivot", "сводн")
        if any(m in name for m in markers):
            return True
        return p.startswith("/session/") and not is_open_workbook(p)

    workbook_norm = norm_path(workbook_path)
    last_norm = norm_path(last_result)

    if user_choice:
        choice_l = user_choice.lower().strip()
        for opt in options:
            if opt["id"].lower() == choice_l or opt["label"].lower() == choice_l:
                return resolve_opt(opt)
            if choice_l in opt["label"].lower():
                return resolve_opt(opt)
            if choice_l in opt["sourcePath"].lower() or choice_l in str(opt["sourceSheet"]).lower():
                return resolve_opt(opt)

    if hint:
        hint_l = hint.lower()
        hits = []
        for opt in options:
            blob = (opt["label"] + " " + opt["sourcePath"] + " " + str(opt["sourceSheet"])).lower()
            if hint_l in blob or any(
                tok and tok in blob for tok in hint_l.replace(",", " ").split()
            ):
                hits.append(opt)
        uniq = []
        for h in hits:
            if h not in uniq:
                uniq.append(h)
        if len(uniq) == 1:
            return resolve_opt(uniq[0])

    if len(options) == 1:
        return resolve_opt(options[0])

    # Default: open R7 workbook beats previous KPI/result files.
    # Follow-up questions about raw columns (оклад, город, …) must not land on *_kpi.xlsx.
    if workbook_norm:
        wb_opts = [o for o in options if norm_path(o["sourcePath"]) == workbook_norm]
        if len(wb_opts) == 1:
            return resolve_opt(wb_opts[0])
        if len(wb_opts) > 1:
            return {
                "ok": True,
                "needsUser": True,
                "options": wb_opts,
                "prompt": "На каком листе открытой книги работать?",
            }

    primary = [o for o in options if not is_result_artifact(o["sourcePath"])]
    if len(primary) == 1:
        return resolve_opt(primary[0])
    if len(primary) > 1:
        return {
            "ok": True,
            "needsUser": True,
            "options": primary,
            "prompt": "С какой таблицей работать дальше?",
        }

    # Only result artifacts left — ask (or auto last_result if single).
    if last_norm:
        last_opts = [o for o in options if norm_path(o["sourcePath"]) == last_norm]
        if len(last_opts) == 1:
            return resolve_opt(last_opts[0])

    return {
        "ok": True,
        "needsUser": True,
        "options": options,
        "prompt": "С какой таблицей работать дальше?",
    }
