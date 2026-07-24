async def handler(state, params):
    params = params if isinstance(params, dict) else {}
    vfs = get_vfs(state)
    if not vfs:
        return {"ok": False, "error": "VFS недоступен"}

    source_path = get_string(params, "sourcePath").strip()
    if not source_path:
        return {"ok": False, "error": "sourcePath обязателен"}
    sheet_a = get_string(params, "sheetA").strip()
    sheet_b = get_string(params, "sheetB").strip()
    if not sheet_a or not sheet_b:
        return {"ok": False, "error": "sheetA и sheetB обязательны"}
    key_fields = [str(x or "").strip() for x in get_array(params, "keyFields") if str(x or "").strip()]
    mode = get_string(params, "mode").strip().lower() or "diff"
    if mode not in ("diff", "join", "reconcile"):
        mode = "diff"
    value_field = get_string(params, "valueField").strip()
    tolerance = 0.01
    raw_tol = params.get("tolerance")
    if isinstance(raw_tol, (int, float)) and math.isfinite(float(raw_tol)):
        tolerance = abs(float(raw_tol))
    target_path = normalize_result_path(get_string(params, "targetPath").strip(), source_path, "compare")

    try:
        from openpyxl import load_workbook

        buffer = await read_workbook_bytes(vfs, source_path)
        wb = load_workbook(io.BytesIO(buffer), data_only=True)
        if sheet_a not in wb.sheetnames or sheet_b not in wb.sheetnames:
            return {
                "ok": False,
                "error": "Листы не найдены. Есть: " + ", ".join(wb.sheetnames),
                "sheets": list(wb.sheetnames),
            }
        a = sheet_to_records(wb, sheet_a)
        b = sheet_to_records(wb, sheet_b)
        if not key_fields:
            common = [h for h in a["headers"] if h in b["headers"]]
            if not common:
                return {"ok": False, "error": "Нет общих колонок для ключа; укажите keyFields"}
            key_fields = [common[0]]

        for field in key_fields:
            if field not in a["headers"] or field not in b["headers"]:
                return {"ok": False, "error": "keyField отсутствует на одном из листов: " + field}

        def index_rows(parsed):
            out = {}
            for row in parsed["rows"]:
                key = "\u0001".join(cell_key(row.get(f)) for f in key_fields)
                if key not in out:
                    out[key] = row
            return out

        map_a = index_rows(a)
        map_b = index_rows(b)
        keys_a = set(map_a.keys())
        keys_b = set(map_b.keys())
        only_a = sorted(keys_a - keys_b)
        only_b = sorted(keys_b - keys_a)
        both = sorted(keys_a & keys_b)

        def row_aoa(headers, rows_list):
            return records_to_aoa(headers, rows_list)

        sheets = []
        stats = {
            "onlyA": len(only_a),
            "onlyB": len(only_b),
            "matched": len(both),
            "changed": 0,
            "reconcileDiffs": 0,
        }

        if mode == "join":
            headers = list(key_fields)
            for h in a["headers"]:
                if h not in headers:
                    headers.append(h + "_A")
            for h in b["headers"]:
                if h not in key_fields:
                    headers.append(h + "_B")
            joined = []
            for key in sorted(keys_a | keys_b):
                ra = map_a.get(key)
                rb = map_b.get(key)
                item = {}
                for f in key_fields:
                    item[f] = (ra or rb).get(f) if (ra or rb) else ""
                if ra:
                    for h in a["headers"]:
                        if h not in key_fields:
                            item[h + "_A"] = ra.get(h)
                if rb:
                    for h in b["headers"]:
                        if h not in key_fields:
                            item[h + "_B"] = rb.get(h)
                joined.append(item)
            sheets.append({"name": "Join", "aoa": row_aoa(headers, joined)})
        else:
            sheets.append({
                "name": "Только_" + sheet_a[:20],
                "aoa": row_aoa(a["headers"], [map_a[k] for k in only_a]),
            })
            sheets.append({
                "name": "Только_" + sheet_b[:20],
                "aoa": row_aoa(b["headers"], [map_b[k] for k in only_b]),
            })
            same_headers = [h for h in a["headers"] if h in b["headers"]]
            matched_rows = []
            changed_rows = []
            for key in both:
                ra = map_a[key]
                rb = map_b[key]
                diff_cols = []
                for h in same_headers:
                    if h in key_fields:
                        continue
                    if cell_key(ra.get(h)) != cell_key(rb.get(h)):
                        diff_cols.append(h)
                if diff_cols:
                    stats["changed"] += 1
                    item = {f: ra.get(f) for f in key_fields}
                    item["changedFields"] = ", ".join(diff_cols)
                    for h in diff_cols:
                        item[h + "_A"] = ra.get(h)
                        item[h + "_B"] = rb.get(h)
                    changed_rows.append(item)
                else:
                    matched_rows.append(ra)
            sheets.append({"name": "Совпадают", "aoa": row_aoa(a["headers"], matched_rows)})
            if changed_rows:
                ch_headers = key_fields + ["changedFields"]
                for row in changed_rows:
                    for k in row.keys():
                        if k not in ch_headers:
                            ch_headers.append(k)
                sheets.append({"name": "Расхождения", "aoa": row_aoa(ch_headers, changed_rows)})

            if mode == "reconcile":
                if not value_field:
                    return {"ok": False, "error": "для mode=reconcile укажите valueField"}
                if value_field not in a["headers"] or value_field not in b["headers"]:
                    return {"ok": False, "error": "valueField нет на обоих листах"}
                rec_rows = []
                for key in both:
                    ra = map_a[key]
                    rb = map_b[key]
                    va = to_number(ra.get(value_field))
                    vb = to_number(rb.get(value_field))
                    if not math.isfinite(va):
                        va = 0.0
                    if not math.isfinite(vb):
                        vb = 0.0
                    delta = va - vb
                    if abs(delta) > tolerance:
                        stats["reconcileDiffs"] += 1
                        item = {f: ra.get(f) for f in key_fields}
                        item[value_field + "_A"] = round_smart(va)
                        item[value_field + "_B"] = round_smart(vb)
                        item["delta"] = round_smart(delta)
                        rec_rows.append(item)
                rec_headers = key_fields + [value_field + "_A", value_field + "_B", "delta"]
                sheets.append({"name": "Сверка_сумм", "aoa": row_aoa(rec_headers, rec_rows)})

        write_info = await write_aoa_workbook(vfs, target_path, sheets, None)
        preview = sheets[0]["aoa"][: min(8, len(sheets[0]["aoa"]))] if sheets else []
        return analytics_done(
            target_path,
            "Сверка листов «" + sheet_a + "» / «" + sheet_b + "» (" + mode + ").",
            sourcePath=source_path,
            sheetA=sheet_a,
            sheetB=sheet_b,
            keyFields=key_fields,
            mode=mode,
            stats=stats,
            resultSheets=[s["name"] for s in sheets],
            preview=preview,
            **write_info,
        )
    except Exception as err:
        return {"ok": False, "error": str(err), "sourcePath": source_path}
